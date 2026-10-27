#!/usr/bin/env python3
"""Seed SyncRove Garmin sessions; do not download study data or contact REDCap.

Run this script from any folder:
    python seed_syncrove_tokens.py
    python seed_syncrove_tokens.py --auth garth
    python seed_syncrove_tokens.py --file "accounts.xlsx" --sheet "Sheet1"
    python seed_syncrove_tokens.py --file "accounts.xlsx" --dry-run

The workbook picker opens when --file is omitted. A worksheet containing
username/password is selected automatically if there is exactly one match.
Only credentials and optional Study ID are read; study dates do not limit seeding.

Authentication follows the installed garminconnect version (or --auth):
  garth:  >=0.2.40,<0.3.0, normally in your existing .venv-old.
  modern: >=0.3.10,<0.4.0; development checked against 0.3.15.

IMPORTANT: Garth's maintainer reports fresh login is broken/deprecated:
https://github.com/matin/garth
Legacy support is provided for the environment you reported working; it is not
a guarantee that Garmin will accept new logins. The modern backend is the
maintained path. This script never changes or upgrades your environment.

Storage defaults to the folder containing the SELECTED ACCOUNT WORKBOOK,
including a mounted remote drive. The script's location does not select storage:
  .garmin_tokens_v2/<sha256(stripped_lowercase_username)[:24]>/oauth1_token.json
  .garmin_tokens_v2/<same_hash>/oauth2_token.json
  .garmin_tokens_v3/<same_hash>.json  (native DI OAuth JSON)
Override the destination folder with --project-dir. Each backend is kept separate.
--existing-token-dir can name a compatible cache root to copy from; original
files are retained. Destination caches are checked first. To migrate the already
created local sessions to the workbook's remote folder without contacting Garmin:
    python seed_syncrove_tokens.py --file "accounts.xlsx" --auth modern \
        --existing-token-dir "/path/to/old/Seed/.garmin_tokens_v3" \
        --copy-existing-only
This mode copies missing destination caches, keeps existing complete destination
caches, and reports missing source caches instead of attempting authentication.
It cannot be combined with --verify-existing or --reseed-existing.

Existing caches are checked locally and reported as UNVERIFIED by default.
--verify-existing performs token-only login/profile validation and preserves the
old files on any failure. No cache is called live-valid based only on existence.
--reseed-existing deliberately requests fresh sessions for the selected backend;
failed attempts still preserve the previous cache. Other-backend tokens are not
converted, overwritten, or counted as usable for the selected backend.

Every unique account is considered, bottom row first. The bottommost nonblank
password wins if duplicate rows disagree; conflicts are recorded. No intentional
inter-account sleep, retry wave, login cap, or 429 circuit breaker is added.
The Garmin library can make multiple requests/retries inside a login attempt.
MFA accounts are recorded and skipped unless --prompt-mfa is supplied.

Seeding_Results/<UTC run id>/ contains an incrementally saved account CSV and a
JSON run summary. No passwords, email addresses, or token contents are logged.
The token and results folders receive their own .gitignore for new local files.
Daily activity/sleep endpoints are not called. Library login may read account
profile/settings. Neither workbook contents nor REDCap records are modified.

Exit codes: 0 = finished without unresolved accounts; 1 = account failures or
missing credentials; 2 = setup error; 130 = interrupted (progress is saved).
--dry-run performs no network calls and writes no files, even when combined with
--verify-existing or --reseed-existing. --copy-existing-only also makes no network
calls, but writes copied caches and reports. Both modes only require openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import logging
import os
import re
import shutil
import sys
import tempfile
import time
import uuid
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path


USERNAME_HEADERS = ("username", "user name", "watch email", "account email", "garmin email", "email")
PASSWORD_HEADERS = ("password", "watch password", "account password", "garmin password")
STUDY_HEADERS = ("study id", "redcap id", "record id", "id")
TOKEN_FOLDERS = {"garth": ".garmin_tokens_v2", "modern": ".garmin_tokens_v3"}
CSV_FIELDS = ("project", "account_hash", "study_ids", "excel_rows", "auth_backend",
              "status", "elapsed_seconds", "error_type", "http_status", "note",
              "duplicate_password_conflict", "other_backend_cache_present", "token_path")


class MFARequiredError(RuntimeError):
    pass


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def canon(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").lower())).strip()


def account_hash(username):
    return hashlib.sha256(username.strip().lower().encode("utf-8")).hexdigest()[:24]


def cell_text(cell, preserve_spaces=False):
    value = cell.value
    if value is None or cell.data_type == "e":
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and not isinstance(value, bool) and re.fullmatch(r"0+", cell.number_format or ""):
        text = str(value).zfill(len(cell.number_format))
    else:
        text = str(value)
    return text if preserve_spaces else text.strip()


@dataclass
class Account:
    username: str = field(repr=False)
    password: str = field(default="", repr=False)
    rows: list = field(default_factory=list)
    study_ids: list = field(default_factory=list)
    password_conflict: bool = False

    @property
    def key(self):
        return account_hash(self.username)


def find_headers(ws):
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)), 1):
        lookup = {}
        for idx, cell in enumerate(row):
            lookup.setdefault(canon(cell.value), []).append(idx)
        def match(aliases):
            for alias in aliases:
                matches = lookup.get(alias, [])
                if len(matches) > 1:
                    raise ValueError(f"Duplicate '{alias}' columns in worksheet {ws.title!r}.")
                if matches:
                    return matches[0]
            return None
        username = match(USERNAME_HEADERS)
        password = match(PASSWORD_HEADERS)
        if username is not None and password is not None:
            return row_number, username, password, match(STUDY_HEADERS)
    return None


def choose_file(value):
    if value:
        return Path(value).expanduser().resolve()
    selected = ""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        try:
            selected = filedialog.askopenfilename(
                title="Select SyncRove account workbook",
                filetypes=[("Excel workbooks", "*.xlsx *.xlsm")])
        finally:
            root.destroy()
    except Exception:
        pass
    if not selected:
        selected = input("SyncRove workbook path (Enter to cancel): ").strip().strip('"')
    if not selected:
        raise ValueError("No workbook selected.")
    return Path(selected).expanduser().resolve()


def load_accounts(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl in this environment: python -m pip install openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet not in wb.sheetnames:
            raise ValueError(f"Worksheet {sheet!r} not found. Available: {wb.sheetnames}")
        candidates = []
        for name in ([sheet] if sheet else wb.sheetnames):
            headers = find_headers(wb[name])
            if headers:
                candidates.append((name, headers))
        if not candidates:
            raise ValueError("No username/password header row found in the first 30 rows. "
                             "Save Excel first if credentials are formulas; cached values are used.")
        if len(candidates) > 1:
            for i, (name, _) in enumerate(candidates, 1):
                print(f"  {i}. {name}")
            choice = input("Which SyncRove worksheet? Enter its number: ").strip()
            if not choice.isdigit() or not 1 <= int(choice) <= len(candidates):
                raise ValueError("Select one worksheet explicitly with --sheet.")
            candidates = [candidates[int(choice) - 1]]
        name, (header, user_col, pass_col, study_col) = candidates[0]
        rows = list(enumerate(wb[name].iter_rows(min_row=header + 1), header + 1))
        accounts = {}
        missing_username_rows = []
        for row_number, row in reversed(rows):
            username = cell_text(row[user_col])
            password = cell_text(row[pass_col], preserve_spaces=True)
            study_id = cell_text(row[study_col]) if study_col is not None else ""
            if not username:
                if password.strip() or study_id:
                    missing_username_rows.append(row_number)
                continue
            key = username.lower()
            account = accounts.setdefault(key, Account(username))
            account.rows.append(row_number)
            if study_id and study_id not in account.study_ids:
                account.study_ids.append(study_id)
            if password.strip():
                if account.password and account.password != password:
                    account.password_conflict = True
                if not account.password:
                    account.password = password
        return list(accounts.values()), {
            "sheet": name, "header_row": header,
            "unique_accounts": len(accounts),
            "duplicate_rows": sum(len(a.rows) - 1 for a in accounts.values()),
            "missing_username_rows": missing_username_rows,
            "credential_values": "Excel cached values; workbook never saved by this program",
        }
    finally:
        wb.close()


def installed_version():
    try:
        return importlib.metadata.version("garminconnect")
    except importlib.metadata.PackageNotFoundError:
        return None


def select_backend(requested, version, dry_run=False):
    parts = tuple(int(x) for x in re.findall(r"\d+", version or "")[:3])
    detected = ("garth" if (0, 2, 40) <= parts < (0, 3, 0) else
                "modern" if (0, 3, 10) <= parts < (0, 4, 0) else None)
    selected = (detected or "modern") if requested == "auto" else requested
    if not dry_run and (detected is None or detected != selected):
        raise ValueError(f"Installed garminconnect: {version or 'not installed'}; selected backend: {selected}. "
                         "Use >=0.2.40,<0.3.0 for garth or >=0.3.10,<0.4.0 for modern. "
                         "Keep different versions in separate virtual environments.")
    return selected


def token_path(root, key, backend):
    return root / key if backend == "garth" else root / (key + ".json")


def local_cache_ok(path, backend):
    """Structural check only; never described as live authentication success."""
    try:
        if backend == "garth":
            one = json.loads((path / "oauth1_token.json").read_text(encoding="utf-8"))
            two = json.loads((path / "oauth2_token.json").read_text(encoding="utf-8"))
            if not isinstance(one, dict) or not isinstance(two, dict):
                return False
            return (all(isinstance(one.get(k), str) and one[k] for k in ("oauth_token", "oauth_token_secret"))
                    and all(isinstance(two.get(k), str) and two[k] for k in ("access_token", "refresh_token"))
                    and all(k in two for k in ("scope", "jti", "token_type", "expires_in", "expires_at",
                                               "refresh_token_expires_in", "refresh_token_expires_at")))
        data = json.loads(path.read_text(encoding="utf-8"))
        return isinstance(data, dict) and all(isinstance(data.get(k), str) and data[k]
                    for k in ("di_token", "di_refresh_token", "di_client_id"))
    except (OSError, ValueError, TypeError):
        return False


def find_cache(roots, key, backend):
    for root in roots:
        paths = [token_path(root, key, backend)]
        if backend == "modern":
            paths.append(root / key / "garmin_tokens.json")
        for path in paths:
            if local_cache_ok(path, backend):
                return path
    return None


@contextmanager
def isolated_auth_environment():
    """Prevent a global token environment variable selecting another account."""
    names = ("GARMINTOKENS", "GARTH_HOME", "GARTH_TOKEN")
    saved = {name: os.environ.pop(name, None) for name in names}
    try:
        yield
    finally:
        for name, value in saved.items():
            if value is not None:
                os.environ[name] = value
            else:
                os.environ.pop(name, None)


class GarminBackend:
    def __init__(self, kind, prompt_mfa=False):
        self.kind = kind
        self.prompt_mfa = prompt_mfa
        for name in ("garminconnect", "garth", "urllib3", "curl_cffi"):
            logging.getLogger(name).setLevel(logging.CRITICAL)
        from garminconnect import Garmin
        self.Garmin = Garmin

    def mfa(self):
        if not self.prompt_mfa:
            raise MFARequiredError("MFA requires --prompt-mfa.")
        code = input("  Enter the Garmin MFA code (blank to skip account): ").strip()
        if not code:
            raise MFARequiredError("MFA code not supplied.")
        return code

    def save_session(self, account, destination, source=None):
        # Token-only verification supplies NO credentials: a restore failure must
        # not silently become a fresh credential login.
        client = self.Garmin(prompt_mfa=self.mfa) if source else self.Garmin(
            email=account.username, password=account.password, prompt_mfa=self.mfa)
        if source:
            copy_cache(source, destination, self.kind)
            result = client.login(str(destination))
        elif self.kind == "garth":
            result = client.login()  # Garth cannot load a nonexistent token path.
        else:
            result = client.login(str(destination))
        if isinstance(result, tuple) and result and result[0] == "needs_mfa":
            raise MFARequiredError("MFA not completed.")
        if self.kind == "garth":
            client.garth.dump(str(destination))
        else:
            # Explicit dump also catches persistence failures suppressed by login().
            client.client.dump(str(destination))
        if not local_cache_ok(destination, self.kind):
            raise RuntimeError("Authentication returned without complete reusable token files.")


def copy_cache(source, destination, backend):
    if backend == "garth":
        destination.mkdir(mode=0o700, parents=True, exist_ok=True)
        for name in ("oauth1_token.json", "oauth2_token.json"):
            shutil.copyfile(source / name, destination / name)
    else:
        destination.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
        shutil.copyfile(source, destination)


def commit_cache(staged, destination, backend):
    """Only complete sessions reach the canonical path; restore on rename failure."""
    if not local_cache_ok(staged, backend):
        raise ValueError("Incomplete staged token files.")
    private_paths = [staged] if backend == "modern" else [staged, *staged.iterdir()]
    for path in private_paths:
        try:
            path.chmod(0o700 if path.is_dir() else 0o600)
        except OSError:
            pass
    if backend == "modern":
        os.replace(staged, destination)
    else:
        backup = staged.parent / "previous_session"
        had_previous = destination.exists()
        if had_previous:
            destination.rename(backup)
        try:
            staged.rename(destination)
        except BaseException:
            if had_previous:
                backup.rename(destination)
            raise
        if had_previous:
            shutil.rmtree(backup)
    paths = [destination] if backend == "modern" else [destination, *destination.iterdir()]
    for path in paths:
        try:
            path.chmod(0o700 if path.is_dir() else 0o600)
        except OSError:
            pass


def protected_folder(path):
    path.mkdir(mode=0o700, parents=True, exist_ok=True)
    try:
        path.chmod(0o700)
    except OSError:
        pass
    ignore = path / ".gitignore"
    if not ignore.exists():
        ignore.write_text("*\n!.gitignore\n", encoding="utf-8")


def atomic_text(path, text):
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                     prefix=".write-", delete=False) as handle:
        tmp = Path(handle.name)
        handle.write(text)
    try:
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def write_results(path, rows):
    import io
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=CSV_FIELDS)
    writer.writeheader()
    for row in rows:
        safe = dict(row)
        for key, value in safe.items():
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                safe[key] = "'" + value
        writer.writerow(safe)
    atomic_text(path, buffer.getvalue())


def error_info(exc):
    """Classify without emitting exception text/URLs that can contain secrets."""
    seen, chain = set(), []
    current = exc
    while current is not None and id(current) not in seen:
        seen.add(id(current))
        chain.append(current)
        current = current.__cause__ or current.__context__
    status = None
    message = " ".join(str(e).lower() for e in chain)
    for error in chain:
        candidate = getattr(getattr(error, "response", None), "status_code", None)
        if isinstance(candidate, int):
            status = candidate
            break
    if status is None:
        found = re.search(r"\b(401|403|429|500|502|503|504)\b", message)
        status = int(found.group(1)) if found else None
    if any(isinstance(e, MFARequiredError) for e in chain):
        return "mfa_required", status, "Rerun with --prompt-mfa to supply the code."
    if status == 429 or "too many" in message or "rate limit" in message:
        return "rate_limited", status or 429, "Garmin rate limited this account; continuing to the next account."
    if status == 403:
        return "access_denied", status, "Garmin rejected the request."
    if status == 401 or any("Authentication" in type(e).__name__ for e in chain):
        return "authentication_failed", status, "Garmin authentication failed; credentials or authentication flow may be rejected."
    return "failed", status, "Login/verification failed; previous tokens were retained."


def run_seeding(accounts, workbook, project_dir, kind, version, args, workbook_info,
                backend=None):
    started = time.perf_counter()
    root = project_dir / TOKEN_FOLDERS[kind]
    roots = list(dict.fromkeys([root, workbook.parent / TOKEN_FOLDERS[kind],
                               *(Path(p).expanduser().resolve() for p in args.existing_token_dir)]))
    other = "modern" if kind == "garth" else "garth"
    other_roots = [project_dir / TOKEN_FOLDERS[other], workbook.parent / TOKEN_FOLDERS[other]]
    rows = []
    sources = []
    for account in accounts:
        source = find_cache(roots, account.key, kind)
        sources.append(source)
        rows.append(dict(zip(CSV_FIELDS, ["SyncRove", account.key, " | ".join(account.study_ids),
            " | ".join(map(str, account.rows)), kind, "not_attempted", "", "", "", "",
            account.password_conflict, bool(find_cache(other_roots, account.key, other)),
            str(token_path(root, account.key, kind))])))
    print(f"Project: SyncRove\nProject folder: {project_dir}\nWorkbook: {workbook}")
    print(f"Worksheet: {workbook_info['sheet']}\nAuthentication: {kind} (garminconnect {version or 'not installed'})")
    print(f"Token folder: {root}\nUnique accounts: {len(accounts)}\nOrder: bottom to top")
    print("Added inter-account delay: 0 seconds; account failures do not stop the loop.")
    if kind == "garth":
        print("NOTE: Garth is deprecated; upstream reports new logins are broken. Existing sessions may still work.")
    if args.dry_run:
        for account, source in zip(accounts, sources):
            action = ("keep existing destination cache" if args.copy_existing_only and source == token_path(root, account.key, kind) else
                      "copy existing cache" if args.copy_existing_only and source else
                      "report missing cache (no login)" if args.copy_existing_only else
                      "inspect existing cache" if source and not args.reseed_existing else
                      "fresh login" if account.password else "skip: missing password")
            print(f"  {account.key[:8]}: would {action}")
        print("DRY RUN: no network calls, files, token changes, or live token verification.")
        return 0, None

    protected_folder(root)
    report_root = project_dir / "Seeding_Results"
    protected_folder(report_root)
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ") + "_" + uuid.uuid4().hex[:6]
    report_dir = report_root / run_id
    report_dir.mkdir(mode=0o700)
    report_csv = report_dir / "SyncRove_Seed_Accounts.csv"
    write_results(report_csv, rows)
    summary = {"project": "SyncRove", "run_id": run_id, "started_at": now_iso(),
               "project_dir": str(project_dir), "workbook_path": str(workbook),
               "auth_backend": kind, "garminconnect_version": version,
               "token_root": str(root), "account_key": "sha256(username.strip().lower().encode('utf-8'))[:24]",
               "token_layout": "<hash>/oauth1_token.json + oauth2_token.json" if kind == "garth" else "<hash>.json",
               "worksheet": workbook_info, "verify_existing": args.verify_existing,
               "reseed_existing": args.reseed_existing,
               "copy_existing_only": args.copy_existing_only, "complete": False}
    summary_file = report_dir / "SyncRove_Seed_Summary.json"
    atomic_text(summary_file, json.dumps(summary, indent=2))
    interrupted = False
    try:
        for index, (account, row, source) in enumerate(zip(accounts, rows, sources), 1):
            print(f"[{index}/{len(accounts)}] Account {account.key[:8]} (Excel rows {row['excel_rows']})", flush=True)
            before = time.perf_counter()
            destination = Path(row["token_path"])
            if args.reseed_existing:
                source = None
            row["status"] = "in_progress"
            write_results(report_csv, rows)
            try:
                if source is None and args.copy_existing_only:
                    row["status"] = "missing_cache"
                    row["note"] = "No complete compatible cache found; no login attempted in copy-only mode."
                elif source is None and not account.password:
                    row["status"] = "missing_password"
                    row["note"] = "No cached session or usable cached Excel password; no login attempted."
                elif source == destination and not args.verify_existing:
                    row["status"] = "existing_unverified"
                    row["note"] = "Complete cache present; live validity not checked."
                else:
                    with tempfile.TemporaryDirectory(prefix=".seed-", dir=root) as staging:
                        staged = Path(staging) / ("session" if kind == "garth" else "session.json")
                        if source is not None and not args.verify_existing:
                            copy_cache(source, staged, kind)
                            new_status = "copied_existing_unverified"
                        else:
                            backend.save_session(account, staged, source)
                            new_status = "existing_verified" if source else "seeded"
                        commit_cache(staged, destination, kind)
                    row["status"] = new_status
                    if source is not None and not args.verify_existing:
                        row["note"] = "Copied compatible cache; live validity not checked. Original retained."
            except Exception as exc:
                row["status"], row["http_status"], row["note"] = error_info(exc)
                row["error_type"] = type(exc).__name__
            except KeyboardInterrupt:
                row["status"] = "interrupted"
                raise
            finally:
                row["elapsed_seconds"] = round(time.perf_counter() - before, 4)
                write_results(report_csv, rows)
            print(f"  {row['status']} ({row['elapsed_seconds']:.2f}s)", flush=True)
    except KeyboardInterrupt:
        interrupted = True
    finally:
        counts = dict(Counter(row["status"] for row in rows))
        summary.update(ended_at=now_iso(), complete=not interrupted, interrupted=interrupted,
                       elapsed_seconds=round(time.perf_counter() - started, 4), counts=counts,
                       complete_local_caches=sum(local_cache_ok(token_path(root, a.key, kind), kind) for a in accounts))
        atomic_text(summary_file, json.dumps(summary, indent=2))
        print("\nSEEDING SUMMARY")
        for status, count in sorted(counts.items()):
            print(f"  {status}: {count}")
        print(f"  Missing-username rows: {len(workbook_info['missing_username_rows'])}")
        print(f"  Elapsed: {summary['elapsed_seconds']:.1f}s\nResults: {report_dir}")
    successful = {"seeded", "existing_verified", "existing_unverified", "copied_existing_unverified"}
    unresolved = any(row["status"] not in successful for row in rows) or bool(workbook_info["missing_username_rows"])
    return (130 if interrupted else 1 if unresolved else 0), report_dir


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Seed SyncRove Garmin accounts without downloading study data or uploading to REDCap.")
    parser.add_argument("--file", "-f", help="Excel workbook; omitted opens a file picker.")
    parser.add_argument("--sheet", "-s", help="Worksheet; automatically detected when exactly one matches.")
    parser.add_argument("--project-dir", default=None, help="Token/results destination; default is the selected workbook's folder (including a remote drive).")
    parser.add_argument("--auth", choices=("auto", "garth", "modern"), default="auto", help="Use installed backend, or require a specific one.")
    parser.add_argument("--existing-token-dir", action="append", default=[], help="Additional compatible token root to search; may be repeated.")
    parser.add_argument("--copy-existing-only", action="store_true", help="Copy available caches to the destination and report missing caches; never contact Garmin.")
    parser.add_argument("--dry-run", action="store_true", help="Inspect workbook/caches without network calls or file writes.")
    parser.add_argument("--verify-existing", action="store_true", help="Token-only verification; a failure retains the existing cache.")
    parser.add_argument("--reseed-existing", action="store_true", help="Request fresh sessions even where selected-backend caches exist; preserve caches on failure.")
    parser.add_argument("--prompt-mfa", action="store_true", help="Prompt for MFA; otherwise mark MFA-required accounts and continue.")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    try:
        if args.copy_existing_only and (args.verify_existing or args.reseed_existing):
            raise ValueError("--copy-existing-only cannot be combined with --verify-existing or --reseed-existing.")
        workbook = choose_file(args.file)
        if not workbook.is_file():
            raise ValueError(f"Workbook not found: {workbook}")
        project_dir = Path(args.project_dir).expanduser().resolve() if args.project_dir else workbook.parent
        if not project_dir.is_dir():
            raise ValueError(f"Project folder does not exist: {project_dir}")
        version = installed_version()
        kind = select_backend(args.auth, version, args.dry_run or args.copy_existing_only)
        accounts, info = load_accounts(workbook, args.sheet)
        if not accounts:
            raise ValueError("No account usernames found in the selected worksheet.")
        if args.dry_run or args.copy_existing_only:
            return run_seeding(accounts, workbook, project_dir, kind, version, args, info)[0]
        with isolated_auth_environment():
            backend = GarminBackend(kind, args.prompt_mfa)
            return run_seeding(accounts, workbook, project_dir, kind, version, args, info, backend)[0]
    except KeyboardInterrupt:
        print("Cancelled.")
        return 130
    except (OSError, ValueError, RuntimeError, ImportError) as exc:
        # Setup messages originate locally; account exceptions are classified and
        # redacted separately inside the processing loop.
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())